"""Campaign export planning and ZIP assembly (app.pipelines.export.campaign_archive). Pure."""

import io
import zipfile
from datetime import datetime, timezone

from docx import Document

from app.pipelines.export import campaign_archive as ca

PIECES = [
    {
        "content": "Launch day\nBig news", "platform": "Instagram",
        "publish_scheduled_at": datetime(2026, 10, 3, tzinfo=timezone.utc),
        "media": [{"url": "https://x/a.png", "mime_type": "image/png", "kind": "image"},
                  {"url": "https://x/b.jpg", "mime_type": "image/jpeg", "kind": "image"}],
    },
    {"content": "Launch day\nSecond", "platform": "Instagram", "publish_scheduled_at": "2026-10-03T09:00:00+00:00",
     "media": [{"url": "https://x/c.mp3", "mime_type": "audio/mpeg"}]},
    {"content": "No media here", "platform": "LinkedIn"},
]


def test_media_folder_is_title_date_platform():
    posts = ca.plan_campaign_export(PIECES)
    assert posts[0].media[0].folder == "media/launch-day_2026-10-03_instagram"
    assert posts[0].media[0].name == "image-1.png" and posts[0].media[1].name == "image-2.jpg"


def test_same_title_day_and_platform_do_not_share_a_folder():
    posts = ca.plan_campaign_export(PIECES)
    assert posts[0].media[0].folder != posts[1].media[0].folder


def test_a_post_with_no_date_is_marked_undated():
    assert ca.plan_campaign_export(PIECES)[2].date_label == "undated"


def test_zip_has_the_docx_and_the_media_where_the_docx_says():
    posts = ca.plan_campaign_export(PIECES)
    data = {m.path: b"BYTES" for p in posts for m in p.media}
    zf = zipfile.ZipFile(io.BytesIO(ca.build_zip("Autumn Launch", posts, data)))
    assert "autumn-launch.docx" in zf.namelist()
    text = "\n".join(p.text for p in Document(io.BytesIO(zf.read("autumn-launch.docx"))).paragraphs)
    assert "Big news" in text and "Instagram" in text
    assert "Media: media/launch-day_2026-10-03_instagram/image-1.png" in text
    assert "media/launch-day_2026-10-03_instagram/image-1.png" in zf.namelist()


def test_media_that_could_not_be_downloaded_is_flagged_in_the_docx():
    posts = ca.plan_campaign_export(PIECES)
    zf = zipfile.ZipFile(io.BytesIO(ca.build_zip("C", posts, {})))
    text = "\n".join(p.text for p in Document(io.BytesIO(zf.read("c.docx"))).paragraphs)
    assert "(could not be included)" in text
    assert not any(n.startswith("media/") for n in zf.namelist())


def test_empty_campaign_still_makes_a_valid_zip():
    zf = zipfile.ZipFile(io.BytesIO(ca.build_zip("Empty", [], {})))
    assert zf.namelist() == ["empty.docx"]
