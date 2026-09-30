"""Campaign export: one ZIP holding a .docx of every post's text and a media/ folder.

Layout:
  <campaign-title>.docx         every post (platform, date, text) and which folder holds its media
  media/<title_date_platform>/  that post's images / audio / video in their native format

Pure: plans the files from documents already read and builds bytes the caller already downloaded.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

from docx import Document

from app.pipelines.export import naming


@dataclass
class CampaignMedia:
    folder: str          # media/<title_date_platform>
    name: str            # file name inside the folder
    url: str
    label: str = ""

    @property
    def path(self) -> str:
        return f"{self.folder}/{self.name}"


@dataclass
class CampaignPost:
    title: str
    platform: str
    date_label: str
    text: str
    media: list[CampaignMedia] = field(default_factory=list)


def _date_label(piece: dict) -> str:
    value = piece.get("publish_scheduled_at") or piece.get("created_at")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str) and len(value) >= 10:
        return value[:10]
    return "undated"


def plan_campaign_export(pieces: list[dict]) -> list[CampaignPost]:
    """One CampaignPost per piece, in the order given. Media goes in a folder named
    title_date-of-publish_platform; folders and files stay unique."""
    posts: list[CampaignPost] = []
    folder_names: list[str] = []
    for p in pieces:
        title = naming.first_line_title(p.get("content"))
        platform = str(p.get("platform") or "")
        date = _date_label(p)
        folder_names.append(
            "_".join([naming.slugify(title), date, naming.slugify(platform, fallback="general", max_chars=24)])
        )
        posts.append(CampaignPost(title=title, platform=platform, date_label=date, text=str(p.get("content") or "")))

    # Folder names may repeat (same title, day and platform); keep each post's media apart.
    unique = naming.unique_names(folder_names)
    for post, piece, folder in zip(posts, pieces, unique):
        items = [m for m in (piece.get("media") or []) if m.get("url")]
        for i, m in enumerate(items, start=1):
            kind = str(m.get("kind") or (m.get("mime_type") or "file").split("/")[0])
            ext = naming.extension_for(m.get("mime_type"), "bin")
            post.media.append(CampaignMedia(
                folder=f"media/{folder}",
                name=f"{naming.slugify(kind, fallback='file', max_chars=12)}-{i}.{ext}",
                url=m["url"],
                label=post.title,
            ))
    return posts


def build_docx(campaign_title: str, posts: list[CampaignPost], missing: Optional[set[str]] = None) -> bytes:
    """A readable document: one section per post. A media line names the file to look for, and
    says so when it could not be downloaded."""
    missing = missing or set()
    doc = Document()
    doc.add_heading(campaign_title or "Campaign", level=0)
    doc.add_paragraph(f"{len(posts)} post{'s' if len(posts) != 1 else ''}")
    for i, post in enumerate(posts, start=1):
        doc.add_heading(f"{i}. {post.title[:80]}", level=1)
        doc.add_paragraph(f"{post.platform or 'General'}  |  {post.date_label}")
        for line in post.text.split("\n"):
            doc.add_paragraph(line)
        for m in post.media:
            note = " (could not be included)" if m.path in missing else ""
            doc.add_paragraph(f"Media: {m.path}{note}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_zip(campaign_title: str, posts: list[CampaignPost], media_bytes: dict[str, bytes]) -> bytes:
    """`media_bytes` is keyed by the media's path inside the ZIP. Media without bytes is skipped
    and marked as such in the document."""
    all_media = [m for p in posts for m in p.media]
    missing = {m.path for m in all_media if m.path not in media_bytes}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{naming.slugify(campaign_title, fallback='campaign')}.docx", build_docx(campaign_title, posts, missing))
        for m in all_media:
            if m.path in media_bytes:
                zf.writestr(m.path, media_bytes[m.path])
    return buf.getvalue()
